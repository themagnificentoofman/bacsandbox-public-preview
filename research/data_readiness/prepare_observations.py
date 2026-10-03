"""Typed observations and descriptive diagnostics, not biological evaluation.

No measurement is dropped, clipped, imputed, or relabeled as detection/extinction.
Means are explicitly over available finite observations, with counts/cells retained.
Published-summary formula checks use only complete groups with no ambiguous zeros.
"""
from collections import Counter, defaultdict
import csv
import hashlib
import json
import math
from pathlib import Path
import re
import statistics

import openpyxl
from research.data_readiness import extract_zhou

HERE = Path(__file__).resolve().parent
OUT = HERE / "results/prepared"
PHASES = ["Weaning", "Weaning + 10 days", "Weaning + 4 weeks"]
TREATMENT = "Mirpur-18 + Glucomannan"
CONTROL = "Mirpur-18"


def read_csv(path):
    with Path(path).open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def parse_value(raw):
    if raw == "":
        return None, "missing_source_cell"
    if raw.startswith("#"):
        return None, "source_excel_error"
    try:
        value = float(raw)
    except ValueError:
        return None, "nonnumeric"
    if not math.isfinite(value):
        return None, "nonfinite"
    return value, "observed_finite"


def stable_id(parts):
    encoded = json.dumps(parts, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(encoded.encode()).hexdigest(), encoded


def canonicalize(name, row, source_hash):
    culture = name == "s1c_growth"
    source_cell = row["log10_cell"] if name == "s9a_copri_fecal" else row["value_cell"]
    raw = row["reported_log10"] if name == "s9a_copri_fecal" else row["value"]
    value, state = parse_value(raw)
    sample_parts = ["culture", row["strain"], row["treatment"], row["replicate"], row["hours"]] if culture else ["mouse"] + [row[k] for k in extract_zhou.META] + [row["phase"], row["specimen"]]
    taxon = row["strain"] if culture else row["organism"]
    obs_id, full_key = stable_id(sample_parts + [taxon])
    flags = []
    if culture:
        flags += ["normalization_descriptions_differ", "replicate_independence_unknown"]
    else:
        flags += ["phase_to_actual_day_unresolved", "dam_litter_cage_unknown", "no_assay_detection_rule"]
    if state != "observed_finite":
        flags.append(state)
    if value == 0:
        flags.append("reported_normalized_OD_zero" if culture else "log_table_zero_semantics_unknown")
    if value is not None and value < 0:
        flags.append("reported_negative_normalized_OD" if culture else "negative_reported_log")
    return {"observation_id": obs_id, "full_key": full_key, "source_section": name,
            "source_file": row["source_file"], "source_sha256": source_hash,
            "source_sheet": row["source_sheet"], "source_cell": source_cell,
            "raw_value": raw, "numeric_value": value, "measurement_state": state,
            "measurement_scale": "published_normalized_OD600" if culture else "published_log10_genome_units_per_g",
            "specimen": "culture_well" if culture else row["specimen"],
            "taxon_or_strain": taxon, "experiment": row.get("experiment", ""),
            "community": row.get("community", ""), "condition": row["treatment"] if culture else row["diet"],
            "sex": row.get("sex", ""), "mouse_id": row.get("mouse_id", ""),
            "replicate": row.get("replicate", ""), "time_hours": row.get("hours", ""),
            "phase_label": row.get("phase", ""),
            "alternate_reported_linear_value": row["value"] if name == "s9a_copri_fecal" else "",
            "alternate_linear_cell": row["value_cell"] if name == "s9a_copri_fecal" else "",
            "metadata_anchors": row["metadata_anchors"], "flags": json.dumps(flags)}


def describe(rows):
    values = [r["numeric_value"] for r in rows if r["measurement_state"] == "observed_finite"]
    counts = Counter(r["measurement_state"] for r in rows)
    return {"n_slots": len(rows), "n_finite": len(values),
            "n_missing": counts["missing_source_cell"], "n_error": counts["source_excel_error"],
            "n_nonnumeric": counts["nonnumeric"], "n_nonfinite": counts["nonfinite"],
            "n_zero": sum(v == 0 for v in values), "n_negative": sum(v < 0 for v in values),
            "mean_of_available_reported_values": statistics.mean(values) if values else None,
            "sd_of_available_reported_values": statistics.stdev(values) if len(values) > 1 else None,
            "minimum": min(values) if values else None, "maximum": max(values) if values else None,
            "all_source_cells": json.dumps([r["source_cell"] for r in rows]),
            "unusable_source_cells": json.dumps([r["source_cell"] for r in rows if r["measurement_state"] != "observed_finite"]),
            "interpretation": "descriptive finite-value summary; counts disclosed; not imputation or an analysis inclusion decision"}


def summary_tokens(text):
    match = re.fullmatch(r"\s*([+-]?\d+(?:\.\d+)?)\s*±\s*([+-]?\d+(?:\.\d+)?)\s*", text)
    if not match:
        raise ValueError(f"unrecognized summary: {text!r}")
    tokens = match.groups()
    precision = [len(v.split(".")[1]) if "." in v else 0 for v in tokens]
    return [float(v) for v in tokens], [0.5 * 10 ** -p for p in precision]


def check_summary(rows, published_text):
    states = [parse_value(r["value"]) for r in rows]
    result = {"n_slots": len(rows), "n_missing": sum(s == "missing_source_cell" for _, s in states),
              "n_errors": sum(s != "observed_finite" and s != "missing_source_cell" for _, s in states),
              "n_zero": sum(v == 0 for v, _ in states), "published_text": published_text,
              "source_value_cells": json.dumps([r["value_cell"] for r in rows])}
    if not rows or any(s != "observed_finite" or v == 0 for v, s in states):
        return result | {"diagnostic_status": "not_tested_flagged_group", "mean_log_sd_log_matches": None, "log_linear_mean_upper_sd_matches": None}
    values = [v for v, _ in states]
    if len(values) < 2:
        return result | {"diagnostic_status": "not_tested_singleton", "mean_log_sd_log_matches": None, "log_linear_mean_upper_sd_matches": None}
    obs, tolerance = summary_tokens(published_text)
    mean_log = statistics.mean(values)
    sd_log = statistics.stdev(values)
    # Stable log10(mean(10**x)) and upper log displacement of a linear SD.
    offset = max(values)
    linear = [10 ** (x - offset) for x in values]
    mean_linear, sd_linear = statistics.mean(linear), statistics.stdev(linear)
    alt_mean = offset + math.log10(mean_linear)
    alt_spread = math.log10(mean_linear + sd_linear) - math.log10(mean_linear)
    matches = lambda pair: all(abs(a - b) <= t + 1e-10 for a, b, t in zip(obs, pair, tolerance))
    return result | {"diagnostic_status": "tested_complete_zero_free_group",
                     "mean_log": mean_log, "sd_log": sd_log, "log_linear_mean": alt_mean, "upper_log_displacement_of_linear_sd": alt_spread,
                     "mean_log_sd_log_matches": matches([mean_log, sd_log]), "log_linear_mean_upper_sd_matches": matches([alt_mean, alt_spread])}


def write_union_csv(path, rows):
    fields = list(dict.fromkeys(k for row in rows for k in row))
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def build(source_dir=None, output_dir=None):
    source_dir = Path(source_dir or extract_zhou.OUT)
    out = Path(output_dir or OUT)
    out.mkdir(parents=True, exist_ok=True)
    manifest = json.loads(extract_zhou.MANIFEST.read_text())
    hashes = {Path(e["key"]).name: e["sha256"] for e in manifest["files"]}
    for name, digest in hashes.items():
        if hashlib.sha256((extract_zhou.RAW / name).read_bytes()).hexdigest() != digest:
            raise ValueError("original supplement hash mismatch")
    names = ["s1c_growth", "s9a_copri_fecal", "s9b_community", "s9c_community", "s9d_community", "s9e_community"]
    inventory = json.loads((source_dir / "inventory.json").read_text())
    if inventory["source_hashes_verified"] != hashes:
        raise ValueError("extraction source hashes do not match pinned supplements")
    for name in names:
        filename = name + ".csv"
        if hashlib.sha256((source_dir / filename).read_bytes()).hexdigest() != inventory["extracted_csv_hashes"][filename]:
            raise ValueError(f"extracted CSV checksum mismatch: {filename}; rebuild from originals")
    raw = {name: read_csv(source_dir / (name + ".csv")) for name in names}
    observations = [canonicalize(name, row, hashes[row["source_file"]]) for name, rows in raw.items() for row in rows]
    if len({r["observation_id"] for r in observations}) != len(observations):
        raise ValueError("duplicate canonical observation identity")
    write_union_csv(out / "observations.csv", observations)
    groups = defaultdict(list)
    for row in observations:
        if row["source_section"] not in ("s1c_growth", "s9a_copri_fecal"):
            continue
        k = tuple(row[c] for c in ("source_section", "taxon_or_strain", "experiment", "condition", "time_hours", "phase_label"))
        groups[k].append(row)
    summaries = [dict(zip(("source_section", "taxon_or_strain", "experiment", "condition", "time_hours", "phase_label"), k), measurement_scale=rows[0]["measurement_scale"], **describe(rows)) for k, rows in sorted(groups.items())]
    write_union_csv(out / "descriptive_groups.csv", summaries)
    contrasts = []
    for summary in summaries:
        treatment = "1.5% glucomannan" if summary["source_section"] == "s1c_growth" else TREATMENT
        control = "1% glucose" if summary["source_section"] == "s1c_growth" else CONTROL
        if summary["condition"] != treatment:
            continue
        match = next(s for s in summaries if s["condition"] == control and all(s[k] == summary[k] for k in ("source_section", "taxon_or_strain", "experiment", "time_hours", "phase_label")))
        contrasts.append({k: summary[k] for k in ("source_section", "taxon_or_strain", "experiment", "time_hours", "phase_label")} | {
            "treatment": treatment, "control": control, "treatment_n_slots": summary["n_slots"], "treatment_n_finite": summary["n_finite"],
            "control_n_slots": match["n_slots"], "control_n_finite": match["n_finite"],
            "measurement_scale": summary["measurement_scale"],
            "difference_of_available_reported_value_means": summary["mean_of_available_reported_values"] - match["mean_of_available_reported_values"],
            "treatment_cells": summary["all_source_cells"], "control_cells": match["all_source_cells"],
            "interpretation": "descriptive treatment-minus-control; unpaired group means, no significance, no selected endpoint or success rule"})
    write_union_csv(out / "descriptive_contrasts.csv", contrasts)
    sheet = openpyxl.load_workbook(extract_zhou.RAW / "pnas.2521522122.sd09.xlsx", data_only=False)["Dataset S9"]
    diagnostics = []
    for name, start, end, phases in [
        ("s9b_community", 63, 79, [(PHASES[0], 4), (PHASES[1], 7), (PHASES[2], 10)]),
        ("s9c_community", 205, 222, [(PHASES[2], 4)]),
        ("s9d_community", 270, 286, [(PHASES[0], 4), (PHASES[1], 7), (PHASES[2], 10)]),
        ("s9e_community", 409, 425, [(PHASES[2], 4)])]:
        for r in range(start, end + 1):
            taxon = sheet.cell(r, 3).value
            taxon = "S. copri" if taxon == "Segatella copri" else taxon
            for phase, c in phases:
                for diet, col in [(TREATMENT, c), (CONTROL, c + 1)]:
                    rows = [row for row in raw[name] if row["organism"] == taxon and row["phase"] == phase and row["diet"] == diet]
                    diagnostics.append({"section": name, "taxon": taxon, "phase": phase, "diet": diet, "summary_cell": sheet.cell(r, col).coordinate} | check_summary(rows, sheet.cell(r, col).value))
    write_union_csv(out / "published_summary_diagnostics.csv", diagnostics)
    summary_counts = {}
    for name in names[2:]:
        items = [d for d in diagnostics if d["section"] == name]
        summary_counts[name] = {"groups": len(items), "tested_complete_zero_free": sum(d["diagnostic_status"].startswith("tested") for d in items),
                                "mean_log_sd_log_matches": sum(d["mean_log_sd_log_matches"] is True for d in items),
                                "log_linear_mean_upper_sd_matches": sum(d["log_linear_mean_upper_sd_matches"] is True for d in items)}
    result = {"purpose": "descriptive preparation only; biological_evaluation_disabled", "observations": len(observations),
              "states": dict(Counter(r["measurement_state"] for r in observations)), "descriptive_groups": len(summaries),
              "descriptive_contrasts": len(contrasts), "published_summary_checks": summary_counts,
              "rounding_rule": "within half last printed decimal unit plus 1e-10 arithmetic tolerance; not a biological threshold",
              "source_csv_hashes": {n: hashlib.sha256((source_dir / (n + ".csv")).read_bytes()).hexdigest() for n in names}}
    (out / "diagnostics.json").write_text(json.dumps(result, indent=2) + "\n")
    return result


if __name__ == "__main__":
    # Rebuild derived inputs from the pinned originals before preparing observations.
    extract_zhou.build()
    print(json.dumps(build(), indent=2))
