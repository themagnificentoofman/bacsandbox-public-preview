"""Source-preserving ZHOU_BM001 mapping; never emits runnable founder counts.

The derived inverse-log fields are arithmetic, not empirical calibration.
Zero-valued log cells remain unresolved. See DOCUMENTATION.md section 10.6.
"""
import argparse
from collections import Counter
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import tempfile

import openpyxl

from research.data_readiness.extract_zhou import META, metadata, numeric_value

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
COMMUNITY = "Community with S. copri BgF5_2"
PHASES = ("Weaning", "Weaning + 10 days", "Weaning + 4 weeks")
DIETS = ("Mirpur-18 + Glucomannan", "Mirpur-18")
MISSING_KEY = ("Iteration 2", COMMUNITY, DIETS[0], "Male", "14")


def key(row):
    return tuple(str(row[field]) for field in META)


def identifier(parts):
    return hashlib.sha256(json.dumps(list(parts), ensure_ascii=False,
                                     separators=(",", ":")).encode()).hexdigest()


def index_unique(rows, fields):
    result = {}
    for row in rows:
        identity = tuple(str(row[f]) for f in fields)
        if any(row[f] is None or not str(row[f]).strip() for f in fields) or identity in result:
            raise ValueError(f"missing or duplicate key: {identity}")
        result[identity] = row
    return result


def classify_log(value):
    if value is None:
        return "missing_source_cell", None
    if isinstance(value, str) and value.startswith("#"):
        return "source_excel_error", None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return "nonnumeric", None
    if not math.isfinite(value):
        return "nonfinite", None
    if value == 0:
        return "log_zero_semantics_unresolved", None
    try:
        linear = 10.0 ** value
    except OverflowError:
        return "inverse_log_out_of_range", None
    if not math.isfinite(linear) or linear <= 0:
        return "inverse_log_out_of_range", None
    return "finite_nonzero_reported_log", linear


def endpoint_mask(row, baseline):
    observed = row["measurement_state"] == "finite_nonzero_reported_log"
    initial = row["phase_label"] == PHASES[0]
    return {"animal_id": row["animal_id"], **{k: row[k] for k in META},
            "phase_label": row["phase_label"], "source_cell": row["source_cell"],
            "measurement_available": observed,
            "endpoint_eligible": observed and not initial,
            "paired_change_eligible": observed and not initial and
                baseline["measurement_state"] == "finite_nonzero_reported_log",
            "reason": ("baseline_input" if initial else "observed_endpoint" if observed
                       else "missing_endpoint_only" if row["measurement_state"] == "missing_source_cell"
                       else "unresolved_measurement"),
            "simulation_ready": False}


def load_verified(source_dir):
    entries = {Path(e["key"]).name: e for e in
               json.loads((HERE / "provenance/supplements.json").read_text())["files"]}
    books, hashes = {}, {}
    for number in (7, 8, 9):
        name = f"pnas.2521522122.sd{number:02d}.xlsx"
        raw = (source_dir / name).read_bytes()
        entry = entries[name]
        digest = hashlib.sha256(raw).hexdigest()
        if len(raw) != entry["size"] or digest != entry["sha256"]:
            raise ValueError(f"source integrity failure: {name}")
        books[number] = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)[f"Dataset S{number}"]
        hashes[name] = digest
    return books, hashes


def observation(sheet, row, col, meta, anchors, phase, taxon, strain_id, hashes,
                linear_col=None):
    cell = sheet.cell(row, col)
    raw = numeric_value(cell)
    state, derived = classify_log(raw)
    supplied = numeric_value(sheet.cell(row, linear_col)) if linear_col else None
    if linear_col:
        if (raw is None) != (supplied is None):
            raise ValueError(f"inconsistent supplied scales at row {row}")
        if supplied is not None and (isinstance(supplied, bool) or
                not isinstance(supplied, (int, float)) or not math.isfinite(supplied) or
                supplied <= 0 or state != "finite_nonzero_reported_log" or
                not math.isclose(math.log10(supplied), raw, rel_tol=0, abs_tol=1e-12)):
            raise ValueError(f"inconsistent supplied scales at row {row}")
    name = "pnas.2521522122.sd09.xlsx"
    return {"observation_id": identifier(key(meta) + ("feces", phase, taxon)),
            "animal_id": identifier(key(meta)), **{k: str(meta[k]) for k in META},
            "specimen": "feces", "phase_label": phase, "taxon": taxon,
            "strain_id": strain_id, "raw_log10": raw, "measurement_state": state,
            "raw_units": "log10(genome units / g feces)",
            "derived_units": "genome units / g feces",
            "derived_value_status": "arithmetic_inverse_log_not_agent_count" if derived is not None
                else "unavailable_pending_measurement_resolution",
            "derived_genome_units_per_g": derived,
            "supplied_genome_units_per_g": supplied,
            "source_file": name, "source_sha256": hashes[name], "source_sheet": sheet.title,
            "source_cell": cell.coordinate,
            "linear_source_cell": sheet.cell(row, linear_col).coordinate if linear_col else "",
            "metadata_anchors": json.dumps(anchors, sort_keys=True)}


def build(source_dir, species_path):
    sheets, hashes = load_verified(Path(source_dir))
    s7, s8, s9 = (sheets[i] for i in (7, 8, 9))
    if tuple(s9.cell(6, c).value for c in (11, 12, 13)) != PHASES:
        raise ValueError("S9A phase headers changed")
    with Path(species_path).open(newline="") as stream:
        profiles = list(csv.DictReader(stream))
    index_unique(profiles, ["name"])
    names = {r["name"] for r in profiles}
    roster = []
    for r, meta, anchors in metadata(s7, 4, 21,
            {2: "gavage_group", 3: "taxon", 4: "strain", 5: "reference"}):
        label = "S. copri" if meta["taxon"] == "Segatella copri" else meta["taxon"]
        # A label match does not establish parameter/genome equivalence.
        candidate = "Segatella copri BgF5_2 (cultured representative)" if label == "S. copri" else ""
        candidate = candidate if candidate in names else ""
        roster.append({"strain_id": identifier((meta["taxon"], meta["strain"])),
                       "taxon": meta["taxon"], "strain": meta["strain"],
                       "observed_taxon": label, "gavage_group": meta["gavage_group"],
                       "model_profile_candidate": candidate,
                       "mapping_status": "label_candidate_parameters_unvalidated" if candidate
                           else "no_exact_strain_profile",
                       "source_file": "pnas.2521522122.sd07.xlsx",
                       "source_sha256": hashes["pnas.2521522122.sd07.xlsx"],
                       "source_sheet": s7.title, "source_range": f"B{r}:E{r}",
                       "metadata_anchors": json.dumps(anchors, sort_keys=True)})
    taxon_map = {r["observed_taxon"]: r for r in roster}
    if len(taxon_map) != 18 or len({r["strain_id"] for r in roster}) != 18:
        raise ValueError("ambiguous strain roster")
    columns = dict(zip(range(2, 7), META))
    animals = [{**{k: str(v) for k, v in meta.items()}, "source_row": r,
                "source_range": f"B{r}:H{r}", "weight_g": numeric_value(s8.cell(r, 7))}
               for r, meta, _ in metadata(s8, 4, 80, columns)]
    animal_index = index_unique(animals, META)
    targets = []
    for r, meta, anchors in metadata(s9, 7, 45, columns):
        if meta["community"] != COMMUNITY or meta["diet"] not in DIETS:
            raise ValueError("unexpected primary arm")
        if key(meta) not in animal_index:
            raise ValueError("target animal absent from S8")
        for col, phase in zip((11, 12, 13), PHASES):
            targets.append(observation(s9, r, col, meta, anchors, phase, "S. copri",
                          taxon_map["S. copri"]["strain_id"], hashes, linear_col=col-4))
    target_index = index_unique(targets, META + ["phase_label"])
    initial = {key(r): r for r in targets if r["phase_label"] == PHASES[0]}
    founders = [dict(r) for r in targets if r["phase_label"] == PHASES[0]]
    community = {}
    if {s9.cell(82, c).value for c in range(8, 25)} != set(taxon_map) - {"S. copri"}:
        raise ValueError("S9B roster changed")
    for r, meta, anchors in metadata(s9, 83, 199, columns | {7: "phase"}):
        identity = key(meta)
        if identity not in initial or meta["phase"] not in PHASES:
            raise ValueError("unmatched community animal or phase")
        cells = []
        for col in range(8, 25):
            label = s9.cell(82, col).value
            obs = observation(s9, r, col, meta, anchors, meta["phase"], label,
                              taxon_map[label]["strain_id"], hashes)
            cells.append(obs)
            if meta["phase"] == PHASES[0]:
                founders.append(obs)
        ck = identity + (meta["phase"],)
        if ck in community:
            raise ValueError("duplicate community observation")
        community[ck] = cells
    index_unique(founders, META + ["strain_id"])
    cecal = {}
    if s9["H225"].value != "S. copri":
        raise ValueError("cecal target header changed")
    for r, meta, _ in metadata(s9, 226, 264, columns | {7: "phase"}):
        ck = key(meta)
        if ck in cecal:
            raise ValueError("duplicate cecal animal")
        cecal[ck] = {"source_cell": f"H{r}", "raw_log10": numeric_value(s9.cell(r, 8)),
                     "phase_label": meta["phase"], "specimen": "cecal contents"}
    if set(cecal) != set(initial):
        raise ValueError("cecal/fecal animal mismatch")
    masks = [endpoint_mask(r, initial[key(r)]) for r in targets]
    animal_rows = []
    for identity, baseline in initial.items():
        obs = [r for r in founders if key(r) == identity]
        if len(obs) != 18:
            raise ValueError("incomplete baseline vector")
        unresolved = [r for r in obs if r["measurement_state"] != "finite_nonzero_reported_log"]
        animal_rows.append({"animal_id": baseline["animal_id"], **dict(zip(META, identity)),
                           "baseline_slots": len(obs), "unresolved_baseline_slots": len(unresolved),
                           "unresolved_source_cells": json.dumps([r["source_cell"] for r in unresolved]),
                           "composition_values_complete": not unresolved,
                           "agent_counts_available": False, "simulation_ready": False})
    missing = []
    for row in targets:
        if row["measurement_state"] != "missing_source_cell":
            continue
        identity = key(row)
        companions = community[identity + (row["phase_label"],)]
        missing.append({"animal_id": row["animal_id"], "compound_key": list(identity),
                        "phase_label": row["phase_label"], "fecal_target_log_cell": row["source_cell"],
                        "fecal_target_linear_cell": row["linear_source_cell"],
                        "companion_cells": [r["source_cell"] for r in companions],
                        "all_companions_missing": all(r["measurement_state"] == "missing_source_cell"
                                                       for r in companions),
                        "cecal_evidence": cecal[identity],
                        "s8_weight_evidence": {k: animal_index[identity][k]
                                              for k in ("source_range", "weight_g")},
                        "handling": "exclude this fecal endpoint only; retain earlier observations",
                        "numeric_recovery": "not recovered", "cause": "not established"})
    if len(missing) != 1 or tuple(missing[0]["compound_key"]) != MISSING_KEY:
        raise ValueError("missing-endpoint contract changed")
    if not missing[0]["all_companions_missing"]:
        raise ValueError("missing-endpoint companion evidence changed")
    time_rows = []
    for iteration, start in (("Iteration 1", 24), ("Iteration 2", 23)):
        for phase, label, nominal, age in zip(PHASES, ("weaning", "P35", "P53"),
                                            (0, 10, 28), (start, 35, 53)):
            time_rows.append({"experiment": iteration, "phase_label": phase,
                              "analysis_label": label, "reported_postnatal_day": age,
                              "nominal_days_from_phase_label": nominal,
                              "days_from_reported_age_subtraction": age-start,
                              "model_elapsed_seconds": None,
                              "mapping_status": "baseline_origin" if nominal == 0 else
                                  "reported_age_and_phase_discrepancy_unresolved"})
    summary = {"schema_version": 1, "benchmark_id": "ZHOU_BM001",
               "evidence_class": "source_mapping_only; no simulation or calibration",
               "source_hashes": hashes,
               "mapping_script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
               "species_data_sha256": hashlib.sha256(Path(species_path).read_bytes()).hexdigest(),
               "target_slots": len(targets), "animals": len(initial), "baseline_slots": len(founders),
               "target_states": dict(Counter(r["measurement_state"] for r in targets)),
               "baseline_states": dict(Counter(r["measurement_state"] for r in founders)),
               "animals_with_unresolved_baseline_values": sum(not r["composition_values_complete"]
                                                             for r in animal_rows),
               "primary_eligible_by_phase": {p: sum(r["endpoint_eligible"] for r in masks
                                                        if r["phase_label"] == p) for p in PHASES},
               "simulation_ready": False,
               "remaining_requirements": ["baseline zero semantics", "strain parameter provenance",
                    "diet exposure", "agent scale", "observation operator", "model time mapping",
                    "experimental grouping", "meaningful error and calibration protocol"]}
    return {"strain_crosswalk.csv": roster, "baseline_observations.csv": founders,
            "target_observations.csv": targets, "endpoint_eligibility.csv": masks,
            "animal_mapping.csv": animal_rows, "time_crosswalk.csv": time_rows,
            "missing_endpoint.json": missing, "mapping_audit.json": summary}


def write_outputs(payloads, output):
    """Serialize before writing; publish an entirely new directory only."""
    output = Path(output)
    if output.exists():
        raise ValueError("output exists; select a new mapping directory")
    serialized = {}
    for name, value in payloads.items():
        if name.endswith(".json"):
            serialized[name] = json.dumps(value, indent=2, allow_nan=False) + "\n"
        else:
            stream = io.StringIO(newline="")
            writer = csv.DictWriter(stream, fieldnames=list(value[0]))
            writer.writeheader()
            writer.writerows(value)
            serialized[name] = stream.getvalue()
    serialized["checksums.json"] = json.dumps({name: hashlib.sha256(text.encode()).hexdigest()
                                               for name, text in serialized.items()}, indent=2) + "\n"
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".zhou-mapping-", dir=output.parent) as temp:
        staged = Path(temp) / "complete"
        staged.mkdir()
        for name, text in serialized.items():
            (staged / name).write_text(text, encoding="utf-8", newline="")
        # Refuse an output created during serialization, including an empty directory.
        output.mkdir()
        # A filesystem error can leave an incomplete directory; checksums expose it.
        for path in staged.iterdir():
            path.replace(output / path.name)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=HERE / "raw/supplements")
    parser.add_argument("--species", type=Path, default=ROOT / "species_data.csv")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        write_outputs(build(args.source_dir, args.species), args.output)
    except (ValueError, OSError) as exc:
        parser.error(str(exc))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
