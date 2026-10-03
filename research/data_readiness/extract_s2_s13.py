"""Read checksum-pinned S2 identities and S13 chemistry without biological mapping.

All source labels, zero annotations and blanks remain literal. Source rows identify
S2 records; neither whitespace normalization nor likely aliases merge records.
"""
from collections import Counter, defaultdict
from decimal import Decimal
import hashlib
import json
import math
from pathlib import Path
import statistics

import openpyxl

from research.data_readiness.extract_zhou import numeric_value, write_csv

HERE = Path(__file__).resolve().parent
RAW = HERE / 'raw/supplements_s2_s13'
MANIFEST = HERE / 'provenance/supplements_s2_s13.json'
OUT = HERE / 'results/s2_s13'
S2_FIELDS = ['genome', 'genome_type_or_purpose', 'study_or_origin', 'reported_genus', 'reported_species',
             'gtdb_genus_first', 'gtdb_species_first', 'gtdb_version_first',
             'gtdb_genus_second', 'gtdb_species_second', 'gtdb_version_second',
             'bacteroidales_included', 'bifunctional_presence_first_pass', 'pul7_structure',
             'prevotella_segatella_included', 'bifunctional_presence_second_pass',
             'amino_acid_identity_to_bg0019_percent', 'host', 'body_site', 'continent', 'country']
TARGETS = {
    49: ('DSM18205; Prevotella entry', 'Prevotella_copri_DSM18205'),
    271: ('DSM18205; Segatella entry', 'Segatella_copri_DSM18205'),
    544: ('F5_2 internal genome identifier', 'Prevotella_copri_P01C7771108_F5_2_JG_bc1003'),
    561: ('Bg0019; alias from pinned author script line 152', 'Prevotella_copri_P06C888_maxbin.023.mod2'),
    562: ('Bg0018; alias from pinned author script line 146', 'Prevotella_copri_P11C888_bin.008.mod2'),
    2455: ('BgF5_2 label; unresolved relation to row 544', 'Segatella_copri_BgF5_2 '),
}
# name, first/last replicate row, analyte header row, last column, summary first row,
# reported adjusted-P row, replicate unit header, summary unit header
SECTIONS = [('total_monosaccharides', 14, 22, 13, 17, 7, 10, 'D12', 'D5'),
            ('free_monosaccharides', 35, 43, 34, 17, 28, 31, 'D33', 'D26'),
            ('linkages', 56, 64, 55, 47, 49, 52, 'D54', 'D47')]


def verified_books(raw):
    entries = {Path(e['key']).name: e for e in json.loads(MANIFEST.read_text())['files']}
    expected = {'pnas.2521522122.sd02.xlsx', 'pnas.2521522122.sd13.xlsx'}
    if set(entries) != expected:
        raise ValueError('unexpected S2/S13 manifest entries')
    books = {}
    for name, entry in entries.items():
        data = (raw / name).read_bytes()
        if len(data) != entry['size'] or hashlib.sha256(data).hexdigest() != entry['sha256'] or hashlib.md5(data).hexdigest() != entry['md5']:
            raise ValueError(f'source integrity failure: {name}')
        books[name] = openpyxl.load_workbook(raw / name, data_only=False)
    return books, entries


def parse_summary(text):
    parts = text.split('±')
    if len(parts) != 2:
        raise ValueError(f'expected printed mean ± spread: {text!r}')
    decimals = [Decimal(p.strip()) for p in parts]
    if not all(d.is_finite() for d in decimals):
        raise ValueError('nonfinite printed summary')
    return [float(d) for d in decimals], [float(Decimal('0.5') * Decimal(10) ** d.as_tuple().exponent) for d in decimals]


def summarize_source_values(values, printed):
    if len(values) < 2 or any(not isinstance(v, (int, float)) or not math.isfinite(v) for v in values):
        raise ValueError('summary diagnostic needs complete finite source values')
    observed, tolerances = parse_summary(printed)
    calculated = [statistics.mean(values), statistics.stdev(values)]
    matches = [abs(a - b) <= t + 8 * max(math.ulp(a), math.ulp(b)) for a, b, t in zip(observed, calculated, tolerances)]
    return {'calculated_mean': calculated[0], 'calculated_sample_sd': calculated[1],
            'printed_mean_half_unit': tolerances[0], 'printed_sd_half_unit': tolerances[1],
            'mean_matches_printed_precision': matches[0], 'sample_sd_matches_printed_precision': matches[1]}


def merged_header_anchor(sheet, coordinate):
    """Resolve only an explicit merged header range, never a data-cell blank."""
    for area in sheet.merged_cells.ranges:
        if coordinate in area:
            return sheet.cell(area.min_row, area.min_col).coordinate
    return coordinate


def build(raw=None, output=None):
    raw, out = Path(raw or RAW), Path(output or OUT)
    books, entries = verified_books(raw)
    out.mkdir(parents=True, exist_ok=True)
    report = {'source_hashes_verified': {n: e['sha256'] for n, e in entries.items()},
              'purpose': 'source readiness; no model mapping, kinetics, dose conversion or biological evaluation', 'workbooks': {}}
    for name, book in books.items():
        report['workbooks'][name] = [{'sheet': s.title, 'rows': s.max_row, 'columns': s.max_column,
            'formula_cells': sum(c.data_type == 'f' for row in s for c in row),
            'error_cells': sum(c.data_type == 'e' for row in s for c in row)} for s in book]
    name = 'pnas.2521522122.sd02.xlsx'
    sheet = books[name]['Dataset S2']
    if sheet['B3'].value != 'Genome' or sheet['R4'].value != 'Amino Acid Identity (%) to Bg0019 Reference':
        raise ValueError('S2 schema changed')
    schema = {field: {'column': openpyxl.utils.get_column_letter(c),
                     'top_header_cell': sheet.cell(3, c).coordinate, 'top_header_raw': sheet.cell(3, c).value,
                     'subheader_cell': sheet.cell(4, c).coordinate, 'subheader_raw': sheet.cell(4, c).value} for c, field in enumerate(S2_FIELDS, 2)}
    for entry in schema.values():
        anchor = merged_header_anchor(sheet, entry['top_header_cell'])
        entry.update(group_header_cell=anchor, group_header_raw=sheet[anchor].value)
    (out / 's2_columns.json').write_text(json.dumps(schema, indent=2) + '\n')
    genomes, targets, stripped = [], [], defaultdict(list)
    for r in range(5, 2633):
        record = {f: numeric_value(sheet.cell(r, c)) for c, f in enumerate(S2_FIELDS, 2)}
        if not isinstance(record['genome'], str) or not record['genome']:
            raise ValueError(f'missing genome identity B{r}')
        record.update(source_row=r, source_range=f'B{r}:V{r}', source_sheet=sheet.title,
                      source_file=name, source_sha256=entries[name]['sha256'])
        genomes.append(record)
        stripped[record['genome'].strip()].append(r)
        if r in TARGETS:
            label, expected = TARGETS[r]
            if record['genome'] != expected:
                raise ValueError(f'target label mismatch at B{r}')
            targets.append({'target_label': label, **record, 'mapping_status': 'source identity only; model/reference mapping not adopted'})
    write_csv(out / 's2_genomes.csv', genomes)
    write_csv(out / 's2_target_identities.csv', targets)
    collisions = [{'stripped_search_label': key, 'source_rows': json.dumps(rows),
                   'raw_identifiers': json.dumps([sheet.cell(r, 2).value for r in rows]),
                   'disposition': 'search diagnostic only; records not merged'} for key, rows in stripped.items() if len(rows) > 1]
    write_csv(out / 's2_whitespace_collisions.csv', collisions)
    report['s2'] = {'rows': len(genomes), 'target_rows': len(targets),
        'raw_unique_identifiers': len({r['genome'] for r in genomes}),
        'whitespace_modified_search_labels': sum(r['genome'] != r['genome'].strip() for r in genomes),
        'stripped_collision_groups': len(collisions), 'genome_types_literal': dict(Counter(r['genome_type_or_purpose'] for r in genomes)),
        'blank_cells_by_field': {f: sum(r[f] is None for r in genomes) for f in S2_FIELDS},
        'f5_2_annotation_discrepancy': {'row_544_first_pass': sheet['N544'].value, 'row_544_second_pass': sheet['Q544'].value,
            'row_2455_first_pass': sheet['N2455'].value, 'row_2455_second_pass': sheet['Q2455'].value,
            'status': 'separate labels with unresolved sequence equivalence; do not reconcile annotations automatically'},
        'accession_column_present': False}
    name = 'pnas.2521522122.sd13.xlsx'
    sheet = books[name]['Dataset S13']
    observations, summaries, pvalues = [], [], []
    report['s13'] = {}
    for section, first, last, header, lastcol, summary_start, p_row, unit_cell, summary_unit_cell in SECTIONS:
        section_rows = []
        for r in range(first, last + 1):
            if sheet.cell(r, 2).value is None or sheet.cell(r, 3).value not in (1, 2, 3):
                raise ValueError(f'missing/unsupported diet replicate B{r}:C{r}')
            for c in range(4, lastcol + 1):
                record = {'section': section, 'food_input': sheet.cell(r, 2).value, 'replicate_label': sheet.cell(r, 3).value,
                    'analyte_raw': sheet.cell(header, c).value, 'value': numeric_value(sheet.cell(r, c)),
                    'units_raw': sheet[unit_cell].value, 'value_cell': sheet.cell(r, c).coordinate,
                    'diet_cell': f'B{r}', 'replicate_cell': f'C{r}', 'analyte_cell': sheet.cell(header, c).coordinate,
                    'units_cell': unit_cell, 'source_file': name, 'source_sheet': sheet.title, 'source_sha256': entries[name]['sha256'],
                    'limits': 'assay units retained; no mg/g or dose conversion; batch, experiment and replicate independence unknown'}
                observations.append(record); section_rows.append(record)
        for i, start in enumerate(range(first, last + 1, 3)):
            for c in range(4, lastcol + 1):
                vals = [numeric_value(sheet.cell(r, c)) for r in range(start, start + 3)]
                printed = sheet.cell(summary_start + i, c).value
                summaries.append({'section': section, 'food_input': sheet.cell(start, 2).value,
                    'analyte_raw': sheet.cell(header, c).value, 'published_summary_raw': printed,
                    'summary_cell': sheet.cell(summary_start + i, c).coordinate,
                    'source_value_cells': json.dumps([sheet.cell(r, c).coordinate for r in range(start, start + 3)]),
                    'n_reported_replicates': len(vals), 'summary_units_raw': sheet[summary_unit_cell].value,
                    'source_file': name, 'source_sheet': sheet.title, **summarize_source_values(vals, printed)})
        for c in range(4, lastcol + 1):
            pvalues.append({'section': section, 'analyte_raw': sheet.cell(header, c).value,
                'reported_value': numeric_value(sheet.cell(p_row, c)), 'reported_method_label': sheet.cell(p_row, 2).value,
                'value_cell': sheet.cell(p_row, c).coordinate, 'source_file': name, 'source_sheet': sheet.title,
                'disposition': 'published adjusted P-value retained; not recomputed or adopted as a BacSandbox decision rule'})
        vals = [r['value'] for r in section_rows]
        report['s13'][section] = {'records': len(vals), 'analytes': lastcol - 3, 'foods': len({r['food_input'] for r in section_rows}),
            'replicate_labels': sorted({r['replicate_label'] for r in section_rows}), 'units_raw': sheet[unit_cell].value,
            'finite': sum(isinstance(v, (int, float)) and math.isfinite(v) for v in vals),
            'missing': sum(v is None for v in vals), 'zero': sum(v == 0 for v in vals),
            'negative': sum(isinstance(v, (int, float)) and v < 0 for v in vals)}
    write_csv(out / 's13_measurements.csv', observations)
    write_csv(out / 's13_summary_diagnostics.csv', summaries)
    write_csv(out / 's13_published_pvalues.csv', pvalues)
    report['s13_summary_checks'] = {'groups': len(summaries), 'mean_matches': sum(r['mean_matches_printed_precision'] for r in summaries),
        'sample_sd_matches': sum(r['sample_sd_matches_printed_precision'] for r in summaries),
        'rounding_rule': 'half last printed decimal/scientific unit plus 8 ULP arithmetic tolerance; not a biological threshold'}
    report['s13_measurements'] = len(observations)
    report['published_pvalue_cells_retained'] = len(pvalues)
    report['output_hashes'] = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(out.iterdir()) if p.suffix in ('.csv', '.json') and p.name != 'inventory.json'}
    (out / 'inventory.json').write_text(json.dumps(report, indent=2) + '\n')
    return report


if __name__ == '__main__':
    r = build()
    print(json.dumps({k: r[k] for k in ('s2', 's13', 's13_summary_checks', 's13_measurements', 'published_pvalue_cells_retained')}, indent=2))
